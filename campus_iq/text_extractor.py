import os
import pdfplumber
import docx2txt
from pptx import Presentation

# TXT extractor
def extract_text_from_txt(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()

# PDF extractor
def extract_text_from_pdf(file_path):
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text

# DOCX extractor
def extract_text_from_docx(file_path):
    return docx2txt.process(file_path)

# PPT extractor
def extract_text_from_ppt(file_path):
    prs = Presentation(file_path)
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

# Main function to handle all file types
def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".txt":
        content = extract_text_from_txt(file_path)
    elif ext == ".pdf":
        content = extract_text_from_pdf(file_path)
    elif ext == ".docx":
        content = extract_text_from_docx(file_path)
    elif ext == ".pptx":
        content = extract_text_from_ppt(file_path)
    else:
        raise ValueError("Unsupported file type")

    return {
        "file_id": "auto_generated_id",   # backend में actual ID generate होगी
        "filename": os.path.basename(file_path),
        "status": "extracted",
        "content": content
    }
if __name__ == "__main__":
    files = [
        "data/demo_notes.txt",
        "data/demo_syllabus.pdf",
        "uploads/sample.docx",
        "uploads/sample.pptx"
    ]
    for f in files:
        result = extract_text(f)
        print("File:", result["filename"])
        print("Status:", result["status"])
        print("Content Preview:", result["content"][:200])  # सिर्फ पहले 200 characters दिखाओ
        print("-" * 50)
